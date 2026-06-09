"""DOCX loader using python-docx. Splits into sections by Heading styles."""
from langchain_core.documents import Document
from loguru import logger


def docx_loader(file_path: str) -> list[Document]:
    """Load a DOCX file, split into sections by Word Heading styles.

    Each section gets metadata with the correct hierarchy chain
    (title → parent headings → current heading).
    """
    try:
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)

        # Parse paragraphs into (level, text) tuples
        # level 0 = body, level 1-3 = heading
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style_name = para.style.name if para.style else ""
            level = 0
            if style_name.startswith('Heading'):
                try:
                    level = int(style_name.replace('Heading', '').strip())
                except ValueError:
                    level = 1
            paragraphs.append((level, text))

        if not paragraphs:
            logger.warning(f"DOCX file {file_path} produced no text content.")
            return []

        from loader.heading_utils import split_docx_by_headings
        sections = split_docx_by_headings(paragraphs)

        docs = []
        for sec in sections:
            if not sec["text"]:
                continue
            docs.append(Document(page_content=sec["text"], metadata={
                "source": file_path,
                "title": sec["title"],
                "hierarchy": sec["hierarchy"],
                "heading": sec["heading"],
            }))

        logger.info(f"DOCX loaded: {file_path} -> {len(docs)} sections")
        return docs if docs else [Document(page_content='\n'.join(t for _, t in paragraphs), metadata={"source": file_path})]

    except Exception as e:
        logger.error(f"Failed to load DOCX {file_path}: {e}")
        raise
