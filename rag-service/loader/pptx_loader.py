"""PPTX loader using python-pptx. Returns one Document per slide."""
from langchain_core.documents import Document
from loguru import logger


def pptx_loader(file_path: str) -> list[Document]:
    """Load a PPTX file. Returns one Document per slide with slide title as heading.

    Each slide gets its own metadata, so chunks from different slides
    are tagged with the correct slide title by zh_title_enhance.
    """
    try:
        from pptx import Presentation

        prs = Presentation(file_path)
        docs = []

        # Use first slide's title as the overall title
        presentation_title = ""
        if prs.slides and prs.slides[0].shapes.title:
            presentation_title = prs.slides[0].shapes.title.text.strip()

        for slide_num, slide in enumerate(prs.slides):
            parts = []
            slide_title = ""

            # Extract slide title from the title shape
            if slide.shapes.title:
                slide_title = slide.shapes.title.text.strip()

            # Extract all text from shapes
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            parts.append(text)

            content = "\n".join(parts)
            if not content:
                continue

            metadata = {
                "source": file_path,
                "page_number": slide_num + 1,
                "total_slides": len(prs.slides),
            }
            if presentation_title:
                metadata["title"] = presentation_title
            if slide_title:
                metadata["heading"] = slide_title

            docs.append(Document(
                page_content=f"[Slide {slide_num + 1}]\n{content}",
                metadata=metadata,
            ))

        logger.info(f"PPTX loaded: {file_path} -> {len(docs)} slides")
        return docs

    except Exception as e:
        logger.error(f"Failed to load PPTX {file_path}: {e}")
        raise
